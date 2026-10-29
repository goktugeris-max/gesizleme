"""ISO 50001 referans çizgisi (EnRÇ) modellerinin istatistik kalitesini data/iso_model.json'a yazar (ISO 50006: EnB modeli uygunluğu).

Kullanım:
    python collector/iso_model_kalite.py "<TSE GİDEN EVRAK klasörü>"

FRM.06.07 Enerji Referans Çizgileri/<yıl>/<no>.EnRÇ ... .xlsx dosyalarındaki Excel regresyon çıktısından (Regresyon İstatistikleri,
ANOVA, katsayı tablosu) seri başına şunları alır: R², ayarlı R², gözlem sayısı, açıklayıcı değişken sayısı, F anlamlılığı, değişkenlerin
en büyük p-değeri. Regresyon çıktısı olmayan dosyalar "ortalama" yöntemli sayılır (baz dönem ortalaması; model istatistiği yok).
Katsayılar ve tüketim değerleri YAZILMAZ; çıktı yalnız kalite göstergeleridir, bu yüzden şifresiz ve herkese açıktır.
Seri kimlikleri iso50001_to_json.py ile aynıdır (tesis-no). Son yılın dosyası esas alınır.
"""
import datetime
import glob
import json
import os
import re
import sys

import openpyxl

from iso50001_to_json import SERILER

OUT_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "iso_model.json")

# ISO 50006 uygulamasında yaygın eşikler (başlangıç değeri, ekip değiştirebilir)
R2_IYI, R2_ZAYIF, P_ESIK = 0.75, 0.5, 0.05


def sayi(x):
    return float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else None


def bul(ws, etiket):
    for satir in ws.iter_rows():
        for c in satir:
            if isinstance(c.value, str) and c.value.strip() == etiket:
                return c
    return None


def sag(ws, c):
    """Etiketin sağındaki ilk sayı (Excel çıktısı bazen bir sütun kaydırılmış)."""
    for k in range(1, 4):
        v = sayi(ws.cell(row=c.row, column=c.column + k).value)
        if v is not None:
            return v
    return None


def sayfa_istatistik(ws):
    r2 = bul(ws, "R Kare")
    if not r2:
        return None
    out = {"r2": sag(ws, r2)}
    ayarli, gozlem = bul(ws, "Ayarlı R Kare"), bul(ws, "Gözlem")
    out["r2_ayarli"] = sag(ws, ayarli) if ayarli else None
    n = sag(ws, gozlem) if gozlem else None
    out["n"] = int(n) if n is not None else None
    anl = bul(ws, "Anlamlılık F")
    out["f_anlamlilik"] = None
    out["degisken_sayisi"] = None
    if anl:
        alt = ws.cell(row=anl.row + 1, column=anl.column).value
        out["f_anlamlilik"] = sayi(alt)
        df = sayi(ws.cell(row=anl.row + 1, column=anl.column - 4).value)  # df sütunu F'nin 3 sütun solunda; kaymayı da tara
        for k in range(2, 7):
            d = sayi(ws.cell(row=anl.row + 1, column=anl.column - k).value)
            if d is not None and float(d).is_integer() and d <= 10:
                df = d
                break
        out["degisken_sayisi"] = int(df) if df is not None else None
    # katsayı tablosu: "P-değeri" başlığının altı, kesişim hariç
    p = bul(ws, "P-değeri")
    pler = []
    if p:
        r = p.row + 1
        while r < p.row + 12:
            ad = ws.cell(row=r, column=1).value or ws.cell(row=r, column=2).value
            pv = sayi(ws.cell(row=r, column=p.column).value)
            if ad is None and pv is None:
                break
            if isinstance(ad, str) and ad.strip().lower() not in ("kesişim", "kesisim") and pv is not None:
                pler.append(pv)
            r += 1
    out["p_en_buyuk"] = max(pler) if pler else None
    return out


def kalite(st):
    """Tek cümlelik değerlendirme: iyi / orta / zayıf (+ nedeni)."""
    if st["yontem"] != "regresyon":
        return "ortalama", "Model istatistiği yok: baz dönem ortalaması kullanılıyor (değişkenden bağımsız)."
    r2 = st["r2"]
    nedenler = []
    if r2 is None:
        return "belirsiz", "R² okunamadı."
    if st["f_anlamlilik"] is not None and st["f_anlamlilik"] > P_ESIK:
        nedenler.append(f"model anlamlı değil (F anlamlılık {st['f_anlamlilik']:.3f} > {P_ESIK})")
    if st["p_en_buyuk"] is not None and st["p_en_buyuk"] > P_ESIK:
        nedenler.append(f"bir değişken anlamlı değil (p={st['p_en_buyuk']:.3f})")
    if st["n"] is not None and st["degisken_sayisi"] and st["n"] < 10 * st["degisken_sayisi"]:
        nedenler.append(f"gözlem sayısı az ({st['n']})")
    if r2 < R2_ZAYIF:
        return "zayif", f"R²={r2:.2f}: model tüketimin yarısından azını açıklıyor" + ("; " + "; ".join(nedenler) if nedenler else "")
    if r2 < R2_IYI or nedenler:
        return "orta", f"R²={r2:.2f}" + ("; " + "; ".join(nedenler) if nedenler else ": kabul edilebilir ama iyileştirilebilir")
    return "iyi", f"R²={r2:.2f}: model tüketimi iyi açıklıyor"


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    klasor = os.path.join(sys.argv[1], "FRM.06.07 Enerji Referans Çizgileri")
    seriler = {}
    for yil in sorted(d for d in os.listdir(klasor) if d.isdigit()):
        for yol in sorted(glob.glob(os.path.join(klasor, yil, "*.xlsx"))):
            ad = os.path.basename(yol)
            m = re.match(r"(\d+)\.", ad)
            if ad.startswith("~$") or not m or int(m.group(1)) not in SERILER:
                continue
            no = int(m.group(1))
            tesis, bolum, enerji, _ = SERILER[no]
            st = None
            wb = openpyxl.load_workbook(yol, data_only=True)
            for ws in wb:
                st = sayfa_istatistik(ws)
                if st:
                    break
            kayit = {"id": f"{tesis}-{no}", "yil": int(yil), "yontem": "regresyon" if st else "ortalama"}
            if st:
                kayit.update({k: (float(f'{v:.4g}') if isinstance(v, float) else v) for k, v in st.items()})
            kayit["kalite"], kayit["aciklama"] = kalite({**{"r2": None, "n": None, "f_anlamlilik": None, "p_en_buyuk": None, "degisken_sayisi": None}, **kayit})
            kayit["aciklama"] = re.sub(r"(\d)\.(\d)", r"\1,\2", kayit["aciklama"])
            seriler[kayit["id"]] = kayit  # sonraki (daha yeni) yıl öncekinin üstüne yazar
    cikti = {"olusturma": datetime.date.today().isoformat(), "esikler": {"r2_iyi": R2_IYI, "r2_zayif": R2_ZAYIF, "p": P_ESIK},
             "seriler": [seriler[k] for k in sorted(seriler, key=lambda x: (x.split('-')[0], int(x.split('-')[1])))]}
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(cikti, f, ensure_ascii=False, indent=1)
    for s in cikti["seriler"]:
        print(s["id"], s["yil"], s["yontem"], s.get("r2"), s.get("n"), s.get("degisken_sayisi"), s.get("f_anlamlilik"), s.get("p_en_buyuk"), "->", s["kalite"])


if __name__ == "__main__":
    main()
