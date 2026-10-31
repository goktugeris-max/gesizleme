#!/usr/bin/env python3
"""İştirak ve bağlı kuruluşların elektrik tüketimi (Enerji Verimliliği Misyon Ekibi ön analiz Excel'i) → data/istirak.enc.json (ŞİFRELİ).

Kullanım:
  ELEKTRIK_PAROLA="..." python3 collector/istirak_to_json.py "<Istirak_Elektrik_Verimlilik_On_Analiz.xlsx>" [--kontrol] [--acik DOSYA]

Kaynak sayfalar: HamVeri (şirket × tesis × abone × ay, "Analiz Değeri" = düzeltilmiş kWh), Abone Analizi (sözleşme gücü, değerlendirme),
Şirket Özeti (çapraz kontrol), Potansiyel (literatür bantlı tasarruf aralıkları), Kaynaklar. ELEKTRIK_PAROLA elektrik.enc.json ile aynıdır (panel tek parolayla açar).
ALINMAZ: abone numarası (yerine SHA-256 ilk 8 hane), TL tutarları, kaynak hücre adresleri, kişi bilgisi. Potansiyel oranları sahada doğrulanmış değildir (literatür bandı).
Betik, şirket toplamlarını Excel'in Şirket Özeti sayfasıyla karşılaştırır; uyuşmazlık varsa uyarır. --kontrol yalnız doğrulama yapar, dosya yazmaz; --acik şifresiz JSON yazar (yalnız yerel deneme için, depoya koymayın).
"""
import base64
import collections
import hashlib
import json
import os
import re
import sys

import openpyxl

OUT = os.path.join(os.path.dirname(__file__), "..", "data", "istirak.enc.json")
PBKDF2_TUR = 600000
OE = range(1, 10)  # Ocak–Eylül


def sayi(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def sifrele(veri, parola):
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC

    salt, iv = os.urandom(16), os.urandom(12)
    anahtar = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=PBKDF2_TUR).derive(parola.encode("utf-8"))
    sifreli = AESGCM(anahtar).encrypt(iv, json.dumps(veri, ensure_ascii=False, separators=(",", ":")).encode("utf-8"), None)
    b64 = lambda b: base64.b64encode(b).decode("ascii")
    return {"v": 1, "iter": PBKDF2_TUR, "salt": b64(salt), "iv": b64(iv), "data": b64(sifreli)}


def numarasiz(t):
    """Metin içindeki abone/tesisat numaralarını (8-9 haneli) parantezleriyle birlikte atar."""
    return re.sub(r"\s*\(\d{8,9}(?:[,\s/]+\d{8,9})*[^)]*\)", "", str(t or "")).strip()


def tesis_adi(t):
    t = re.sub(r"\s+", " ", numarasiz(t)).strip()
    if t.startswith("[") or t.startswith("(ana listede yok)") or not t:
        return ""
    return t


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    kontrol = "--kontrol" in sys.argv
    acik = sys.argv[sys.argv.index("--acik") + 1] if "--acik" in sys.argv else None
    if acik and acik in args:
        args.remove(acik)
    if not args:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(args[0], data_only=True)

    # ---- HamVeri ----
    ham = list(wb["HamVeri"].iter_rows(min_row=2, values_only=True))
    abone = collections.OrderedDict()
    sirket_ay = collections.defaultdict(lambda: collections.defaultdict(lambda: [None] * 12))
    bulgu = collections.defaultdict(lambda: {"metin": 0, "duzeltme": 0, "dogrulanmadi": 0})
    for r in ham:
        sirket, tesis, no, yil, ay, deger, durum = r[1], r[2], str(r[3]), r[5], r[6], sayi(r[9]), r[10]
        if not sirket or yil is None or ay is None:
            continue
        yil, ay = int(yil), int(ay)
        if durum == "METİN→SAYI":
            bulgu[sirket]["metin"] += 1
        elif durum == "DÜZELTME ÖNERİSİ":
            bulgu[sirket]["duzeltme"] += 1
        elif durum == "DOĞRULANMADI":
            bulgu[sirket]["dogrulanmadi"] += 1
        if deger is None:
            continue
        kimlik = hashlib.sha256(f"{sirket}|{tesis}|{no}".encode()).hexdigest()[:8]
        a = abone.setdefault(kimlik, {"id": kimlik, "sirket": sirket, "tesis": tesis_adi(tesis), "no": no, "tesisHam": str(tesis), "aylik": {}})
        a["aylik"].setdefault(str(yil), [None] * 12)[ay - 1] = round(deger, 2)
        s = sirket_ay[sirket][yil]
        s[ay - 1] = (s[ay - 1] or 0) + deger

    # ---- Abone Analizi: sözleşme gücü, değerlendirme ----
    soz = {}
    for r in wb["Abone Analizi"].iter_rows(min_row=5, values_only=True):
        if r[0] is None or not r[1]:
            continue
        soz[(r[1], str(r[3]))] = {"soz": sayi(r[4]), "not": r[10]}
    sirket_tesis = {}
    for a in abone.values():
        k = (a["sirket"], a["no"])
        if k in soz:
            a["soz"] = soz[k]["soz"]
            a["degerlendirme"] = soz[k]["not"]
        del a["no"], a["tesisHam"]
    for a in abone.values():
        for k in ("soz", "degerlendirme"):
            if a.get(k) is None:
                a.pop(k, None)

    # ---- Şirket Özeti (çapraz kontrol + tür/not) ----
    ozet = {}
    for r in wb["Şirket Özeti"].iter_rows(min_row=5, values_only=True):
        if r[0] and r[0] != "TOPLAM" and r[1]:
            ozet[r[0]] = r
    # ---- Potansiyel ----
    pot = {}
    prm = {}
    for r in wb["Potansiyel"].iter_rows(min_row=1, values_only=True):
        if r[0] and "TEP katsayısı" in str(r[0]):
            prm["tep"] = sayi(r[2])
        if r[0] and "Emisyon faktörü" in str(r[0]):
            prm["ef"] = sayi(r[2])
        if r[0] in ozet and r[1] and sayi(r[3]) is not None:
            pot[r[0]] = {"kategori": r[1], "baz": sayi(r[2]), "dusuk": sayi(r[3]), "yuksek": sayi(r[4]), "onlem": r[13], "kaynak": [x.strip() for x in str(r[14] or "").split(",") if x.strip()]}
    kaynaklar = [{"kod": r[0], "kaynak": r[1], "bilgi": r[2], "url": r[3]} for r in wb["Kaynaklar"].iter_rows(min_row=4, values_only=True) if r[0] and re.match(r"K\d+", str(r[0]))]

    sirketler, uyari = [], []
    for ad, r in ozet.items():
        ay = {str(y): v for y, v in sirket_ay[ad].items()}
        oe = lambda y: sum((sirket_ay[ad][y][i - 1] or 0) for i in OE) if y in sirket_ay[ad] else None
        yillik = lambda y: sum(v or 0 for v in sirket_ay[ad][y]) if y in sirket_ay[ad] else None
        # çapraz kontrol (Excel O–E ve yıllık)
        for ad_, mine, theirs in (("2022 O–E", oe(2022), sayi(r[6])), ("2023 O–E", oe(2023), sayi(r[7])), ("2024 O–E", oe(2024), sayi(r[8])), ("2022 yıllık", yillik(2022), sayi(r[2])), ("2023 yıllık", yillik(2023), sayi(r[3]))):
            if mine is not None and theirs is not None and abs(mine - theirs) > max(1.0, abs(theirs) * 1e-6):
                uyari.append(f"{ad} {ad_}: hesap {mine:,.1f} ≠ Excel {theirs:,.1f}")
        son24 = [i + 1 for i, v in enumerate(sirket_ay[ad].get(2024, [None] * 12)) if v]
        sirketler.append({
            "ad": ad, "tur": r[1], "aylik": ay, "veri24": f"Ocak–{['', 'Ocak', 'Şubat', 'Mart', 'Nisan', 'Mayıs', 'Haziran', 'Temmuz', 'Ağustos', 'Eylül', 'Ekim', 'Kasım', 'Aralık'][max(son24)]}" if son24 else None,
            "ay24": len([i for i in OE if (sirket_ay[ad].get(2024, [None] * 12)[i - 1])]), "kapsam": numarasiz(r[15]) if r[15] else None,
            "abone": sum(1 for a in abone.values() if a["sirket"] == ad), "bulgu": bulgu[ad], "potansiyel": pot.get(ad),
        })

    cikti = {
        "guncelleme": "Ekim 2026 (ön analiz, taslak)", "donem": "2024 Ocak–Eylül", "parametre": prm,
        "kaynak": "Enerji Verimliliği Misyon Ekibi ön analiz dosyası (02.10.2026, taslak); şirketlerin Ekim 2024 tarihli resmi yanıtları",
        "sirketler": sirketler, "aboneler": list(abone.values()), "kaynaklar": kaynaklar,
    }
    toplam = {y: sum((sirket_ay[s["ad"]][y][i - 1] or 0) for s in sirketler for i in OE) for y in (2022, 2023, 2024)}
    print(f"{len(sirketler)} şirket, {len(abone)} abone; O–E toplamları (kWh): " + ", ".join(f"{y}: {v:,.0f}" for y, v in toplam.items()))
    for u in uyari:
        print("UYARI:", u)
    if not uyari:
        print("Şirket toplamları Excel'in Şirket Özeti sayfasıyla tuttu.")
    if kontrol:
        return
    if acik:
        json.dump(cikti, open(acik, "w", encoding="utf-8"), ensure_ascii=False)
        print("Şifresiz yazıldı:", acik)
        return
    parola = os.environ.get("ELEKTRIK_PAROLA", "")
    if len(parola) < 12:
        sys.exit("ELEKTRIK_PAROLA (en az 12 karakter) gerekli; elektrik.enc.json ile aynı olmalı.")
    json.dump(sifrele(cikti, parola), open(OUT, "w"), separators=(",", ":"))
    print("Yazıldı:", os.path.relpath(OUT), os.path.getsize(OUT) // 1024, "KB")


if __name__ == "__main__":
    main()
