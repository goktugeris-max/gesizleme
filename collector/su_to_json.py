#!/usr/bin/env python3
"""Kurumsal su tüketimi özetinden (Yıllık Tüketimler Excel'i) data/kaynak_raporu.json içindeki `su` bölümünü günceller.

Kullanım:
  python3 collector/su_to_json.py "<Yillik_Tuketimler_2026-09.xlsx>" [--kontrol]

Excel düzeni (sayfalar): "2024", "2025", "2026" (AY / TÜKETİM m3 / TUTAR TL), "<yıl> Ayrışık" (Park Bahçe, Mezarlıklar, İşletmeler),
"FATURALANDIRMA ORANI" (abone ve fatura sayıları), "Yüksek Tüketimli Park/İşletme/Merzarlık Ab." (abone bazında 2024-2026 aylık).
YALNIZ m³ ALINIR: TL tutarları, abone numaraları YAZILMAZ (depo herkese açık; yüksek tüketimli abonelerin yalnız tesis adı kalır).
Çıktı `su.aylik`, `su.son_ay`, `su.ayrisik`, `su.fatura`, `su.yuksek_aboneler` alanlarını yazar; dosyanın geri kalanına (hedef, notlar, akaryakıt…) dokunmaz.
"""
import json
import os
import re
import sys

import openpyxl

OUT = os.path.join(os.path.dirname(__file__), "..", "data", "kaynak_raporu.json")
AYLAR = ["ocak", "şubat", "mart", "nisan", "mayıs", "haziran", "temmuz", "ağustos", "eylül", "ekim", "kasım", "aralık"]
HASSAS = re.compile(r"sığınma|siginma|konukevi|kadın konuk|sığınak", re.I)


def kucuk(s):
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def turkce_baslik(ad):
    """BÜYÜK HARFLİ adları Türkçe kurallarıyla 'Baş Harfi Büyük' yapar (str.title() İ/I'yı bozar); karışık yazılı adlara dokunmaz."""
    ad = re.sub(r"\s+", " ", ad).strip()
    if not ad.isupper():
        return ad
    def kelime(w):
        w = w.replace("İ", "i").replace("I", "ı").lower()
        ilk = {"i": "İ", "ı": "I"}.get(w[:1], w[:1].upper())
        return ilk + w[1:]
    return " ".join("-".join(kelime(x) for x in w.split("-")) for w in ad.split(" "))


def sayi(x):
    return float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else None


def aylik_sayfa(ws, sutun):
    """AY başlığından sonraki 12 satırdan `sutun` (0 tabanlı) değerleri; boş ay = None."""
    satirlar = list(ws.iter_rows(values_only=True))
    out = {}
    for r in satirlar:
        ad = kucuk(r[1]) if len(r) > 1 else ""
        if ad in AYLAR and AYLAR.index(ad) not in out:
            out[AYLAR.index(ad)] = sayi(r[sutun])
    return [None if out.get(i) in (None, 0.0) else round(out[i]) for i in range(12)]


def ayrisik_sayfa(ws):
    """Park Bahçe / Mezarlıklar / İşletmeler sütunları (C, D, E)."""
    sut = {"park": 2, "mezarlik": 3, "isletme": 4}
    satirlar = list(ws.iter_rows(values_only=True))
    res = {k: [None] * 12 for k in sut}
    for r in satirlar:
        ad = kucuk(r[0] if r[0] else (r[1] if len(r) > 1 else ""))
        if ad in AYLAR:
            i = AYLAR.index(ad)
            for k, c in sut.items():
                v = sayi(r[c]) if len(r) > c else None
                res[k][i] = None if v in (None, 0.0) else round(v)
    return res


def fatura_sayfasi(ws):
    sut = {"park": (2, 3), "mezarlik": (5, 6), "isletme": (8, 9)}
    res = {k: {"abone": [None] * 12, "fatura": [None] * 12} for k in sut}
    for r in ws.iter_rows(values_only=True):
        ad = kucuk(r[1] if len(r) > 1 else "")
        if ad in AYLAR:
            i = AYLAR.index(ad)
            for k, (a, f) in sut.items():
                va, vf = sayi(r[a]), sayi(r[f])
                res[k]["abone"][i] = None if va in (None, 0.0) else int(va)
                res[k]["fatura"][i] = None if vf in (None, 0.0) else int(vf)
    return res


def yuksek_sayfa(ws):
    """Başlık satırı 'NO AD' ile başlayan bloklar; altında 2024/2025/2026 başlığı ve 12 ay."""
    satirlar = list(ws.iter_rows(values_only=True))
    bloklar, i = [], 0
    while i < len(satirlar):
        r = satirlar[i]
        ilk = next((c for c in r[:3] if c not in (None, "")), None)
        if isinstance(ilk, str) and re.match(r"^\s*\d{4,}\s+\S", ilk):
            ad = turkce_baslik(re.sub(r"^\s*\d{4,}\s+", "", ilk))
            if len(ad) > 60:
                ad = ad[:59].rstrip() + "…"
            yillar = {}
            j = i + 1
            baslik = None
            while j < min(len(satirlar), i + 20):
                rr = satirlar[j]
                vals = [c for c in rr if c is not None]
                if baslik is None and any(isinstance(c, (int, float)) and 2023 < c < 2100 for c in vals):
                    baslik = {int(c): idx for idx, c in enumerate(rr) if isinstance(c, (int, float)) and 2023 < c < 2100}
                elif baslik is not None:
                    ay = next((kucuk(c) for c in rr[:3] if isinstance(c, str) and kucuk(c) in AYLAR), None)
                    if ay is not None:
                        for y, c in baslik.items():
                            v = sayi(rr[c]) if len(rr) > c else None
                            yillar.setdefault(str(y), [None] * 12)[AYLAR.index(ay)] = None if v in (None, 0.0) else round(v)
                    elif any(isinstance(c, str) and kucuk(c) == "toplam" for c in rr[:3]):
                        break
                j += 1
            if yillar and not HASSAS.search(ad):
                bloklar.append({"ad": ad, "aylik": yillar})
            i = j
        i += 1
    return bloklar


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    aylik = {y: aylik_sayfa(wb[y], 2) for y in ("2024", "2025", "2026") if y in wb.sheetnames}
    ayr = {y: ayrisik_sayfa(wb[y + " Ayrışık"]) for y in aylik if y + " Ayrışık" in wb.sheetnames}
    fatura = fatura_sayfasi(wb["FATURALANDIRMA ORANI"]) if "FATURALANDIRMA ORANI" in wb.sheetnames else None
    yuksek = {}
    for ad, anahtar in (("park", "Yüksek Tüketimli Park"), ("isletme", "Yüksek Tüketimli İşletme"), ("mezarlik", "Yüksek Tüketimli Merzarlık")):
        sayfa = next((n for n in wb.sheetnames if n.startswith(anahtar)), None)
        if sayfa:
            yuksek[ad] = yuksek_sayfa(wb[sayfa])
    son_ay = max(i + 1 for i, v in enumerate(aylik["2026"]) if v is not None)
    print("aylık:", {y: [None if v is None else round(v / 1000) for v in a] for y, a in aylik.items()})
    print("son ay:", son_ay, "| yüksek tüketimli:", {k: [b["ad"] for b in v] for k, v in yuksek.items()})
    if "--kontrol" in sys.argv:
        return
    with open(OUT, encoding="utf-8") as f:
        K = json.load(f)
    su = K["su"]
    su["aylik"] = aylik
    su["son_ay"] = son_ay
    su["ayrisik"] = ayr
    if fatura:
        su["fatura"] = fatura
    su["yuksek_aboneler"] = yuksek
    su["kaynak_ozet"] = "Yıllık Tüketimler özeti (Eylül 2026); yalnız m³, TL ve abone numarası yazılmaz"
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(K, f, ensure_ascii=False, indent=1)
        f.write("\n")
    print("yazıldı:", OUT)


if __name__ == "__main__":
    main()
