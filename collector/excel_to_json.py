"""Santral üretim ve arıza Excel dosyalarından data/performans.json üretir.

Kullanım:
    python collector/excel_to_json.py "2025-2026_Ges_Uretim_Bilgileri.xlsx" "GES_ARIZA_TAKIP.xlsx"

Depo ve site herkese açık olduğu için tesisat no, adres ve ihale bedeli GİBİ alanlar
çıktıya YAZILMAZ; yalnızca üretim/performans verisi yazılır.
"""
import json
import os
import re
import sys
from datetime import datetime

import openpyxl

OUT_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "performans.json")
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]

# sayfa öneki -> (kimlik, görünen ad, canlı veri eşleşmesi, "GES Üretim İstatistikleri" sayfasındaki ad)
# canlı eşleşme: panel/durum.json'daki portal ve santral adı (yoksa None)
SANTRALLER = {
    "1-": ("seyrek-hayvan-hastanesi", "Seyrek Hayvan Hastanesi", ("SMA", "ibbseyrekges"), "Seyrek Hayvan Hastanesi"),
    "2-": ("buca-toros-itfaiye", "Buca Toros İtfaiye", ("Huawei", "TR-Buca İtfaiye Ges"), "Buca Toros İtfaiye"),
    "3-": ("ekrem-akurgal", "Ekrem Akurgal Rekreasyon Alanı", None, "Ekrem Akurgal"),
    "4-": ("ornekkoy-sosyal-projeler", "Örnekköy Sosyal Projeler", ("Huawei", "TR-Karşıyaka Örnekköy"), "Karşıyaka Örnekköy"),
    "5-": ("bornova-havuz", "Bornova Yüzme Havuzu", ("GoodWe", "Bornova havuz Ges"), "Bornova Yüzme Havuzu"),
    "6-": ("selcuk-transfer", "Selçuk Katı Atık Transfer İstasyonu", None, "Selçuk Katı Atık Transfer İstasyonu"),
    "7-": ("yesilyurt-kultur-merkezi", "Yeşilyurt (Mustafa Necati) Kültür Merkezi", ("SMA", "yesilyurt kultur merkezi"), "Yeşilyurt Kültür Merkezi"),
    "8-": ("bergama-mezbahasi", "Bergama Mezbahası", ("GoodWe", "Bergama Mezbahası"), "Bergama Mezbahası"),
    "9-": ("aliaga-itfaiye", "Aliağa İtfaiye Merkezi", ("GoodWe", "Aliağa İtfaiye Merkezi"), "Aliağa İtfaiye Merkezi"),
    "10-": ("konak-tuneller", "Konak Tünelleri", ("GoodWe", "Konak tuneller ges"), "Konak Tünelleri"),
    "11-": ("bergama-itfaiye", "Bergama İtfaiye", ("GoodWe", "BERGAMA İTFAİYE"), "Bergama İtfaiye"),
    "12-": ("uzundere-spor", "Uzundere Spor Salonu", ("GoodWe", "Uzundere Can Çağnay Çok Amaçlı Spor Salonu"), "Uzundere Can Çağnay Çok Amaçlı Spor Salonu"),
    "13-": ("aliaga-kultur-merkezi", "Aliağa Kültür Merkezi", ("Huawei", "Aliaga Kultur Merkezi GES"), "Aliağa Kültür Merkezi"),
    "14-": ("seyrek-itfaiye", "Seyrek İtfaiye", ("GoodWe", "SEYREK İTFAİYE"), "Seyrek İtfaiye"),
    "15-": ("bornova-kultur-merkezi", "Bornova Kültür Merkezi", ("Huawei", "Bornova Kultur Merkezi GES"), "Bornova Kültür Merkezi"),
    "16-": ("cigli-aile-danisma", "Çiğli Aile Danışma Merkezi", ("Huawei", "Çiğli Sosyal Hizmet Binası GES"), "Çiğli Aile Danışma Merkezi"),
    "17-": ("gumuldur-itfaiye", "Gümüldür İtfaiye", ("Huawei", "Gumuldur Itfaiye GES"), "Gümüldür İtfaiye"),
    "18-": ("selcuk-itfaiye", "Selçuk İtfaiye", ("Huawei", "Selcuk Itfaiye GES"), "Selçuk İtfaiye"),
    "19-": ("cesme-itfaiye", "Çeşme İtfaiye", ("GoodWe", "Çeşme İtfaiye Ges"), "Çeşme İtfaiye"),
    "20-": ("caybasi-itfaiye", "Çaybaşı İtfaiye", ("GoodWe", "Çaybaşı İtfaiye"), "Çaybaşı İtfaiye"),
    "21-": ("akkaya-dges", "Akkaya Telsiz İstasyonu (Depolamalı)", None, "Akkaya Telsiz İstasyonu Depolamalı"),
}


def sayi(x):
    """Hücre değerini sayıya çevir; sayı değilse None. Virgüllü ondalık ('38,5') desteklenir."""
    if isinstance(x, bool) or x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(str(x).strip().replace(".", "").replace(",", ".")) if "," in str(x) else float(str(x).strip())
    except ValueError:
        return None


def tr_lower(s):
    return str(s).replace("İ", "i").replace("I", "ı").lower()


def tr_title(s):
    return " ".join(w[:1].replace("i", "İ").upper() + tr_lower(w[1:]) if w else w for w in str(s).strip().split(" "))


def yuvarla(x, n=1):
    return None if x is None else round(x, n)


def bul(ws, metin):
    """Hücre değeri `metin`e eşit olan ilk (satır, sütun)."""
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip().lower() == metin.lower():
                return c.row, c.column
    return None, None


def bilgi(ws, anahtar):
    """Sayfanın ilk 3 sütunundaki 'anahtar ...' etiketinin sağındaki değerler (etiket A veya B sütununda olabilir)."""
    for r in range(1, 25):
        for col in (1, 2):
            k = ws.cell(r, col).value
            if isinstance(k, str) and tr_lower(k).strip().startswith(tr_lower(anahtar)):
                return [ws.cell(r, c).value for c in range(col + 1, 4) if ws.cell(r, c).value is not None]
    return []


def santral_sayfasi(ws):
    """Bir santral sayfasından aylık üretim (yıl -> 12 değer), aylık beklenen ve teknik bilgi."""
    r0, c0 = bul(ws, "Ocak")
    if not r0:
        raise ValueError(f"{ws.title}: 'Ocak' satırı bulunamadı")
    baslik = r0 - 1
    # "Beklenen Üretim" sütunu
    beklenen_col = next((c for c in range(1, ws.max_column + 1)
                         if isinstance(ws.cell(baslik, c).value, str) and "beklenen" in ws.cell(baslik, c).value.lower()), None)
    if not beklenen_col:
        raise ValueError(f"{ws.title}: 'Beklenen Üretim' sütunu bulunamadı")

    aylik = {}
    for c in range(c0 + 1, beklenen_col):
        yil = ws.cell(baslik, c).value
        if isinstance(yil, (int, float)) and 2015 <= int(yil) <= 2100:
            aylik[int(yil)] = [sayi(ws.cell(r0 + i, c).value) for i in range(12)]
    beklenen = [sayi(ws.cell(r0 + i, beklenen_col).value) for i in range(12)]

    kap = bilgi(ws, "kapasite")
    akt = bilgi(ws, "aktivasyon")
    marka = [str(m).strip() for m in bilgi(ws, "inverter markas")]
    adet = [sayi(a) for a in bilgi(ws, "inverter sayıs")]
    panel = bilgi(ws, "panel sayıs")
    return {
        "kw": sayi(kap[0]) if kap else None,
        "aktivasyon": akt[0].strftime("%Y-%m-%d") if akt and isinstance(akt[0], datetime) else None,
        "panel_adet": int(sayi(panel[0])) if panel and sayi(panel[0]) is not None else None,
        "inverter": ", ".join(dict.fromkeys(marka)) or None,
        "inverter_adet": int(sum(a for a in adet if a)) if any(adet) else None,
        "aylik": aylik,
        "beklenen_aylik": beklenen,
    }


def uretim_dosyasi(yol):
    wb = openpyxl.load_workbook(yol, data_only=True)

    # Yıllık özet sayfası: kapasite, resmi yıllık toplamlar ve devrede olmayan (planlanan) santraller
    ozet = wb["GES Üretim İstatistikleri"]
    yillar = {c: ozet.cell(2, c).value for c in range(5, 14)}
    ozet_bilgi, planlanan = {}, []
    for r in range(3, ozet.max_row + 1):
        ad, durum, kw = ozet.cell(r, 2).value, ozet.cell(r, 3).value, ozet.cell(r, 4).value
        if not ad or not durum:
            continue
        ad = str(ad).strip()
        if str(durum).strip() == "Devrede":
            ozet_bilgi[ad] = {
                "kw": sayi(kw),
                "yillik": {str(int(y)): yuvarla(sayi(ozet.cell(r, c).value), 1)
                           for c, y in yillar.items() if sayi(ozet.cell(r, c).value) is not None},
            }
        else:
            planlanan.append({"ad": tr_title(ad), "kw": sayi(kw), "durum": str(durum).strip()})

    sonuc, notlar = [], []
    for ws in wb.worksheets:
        m = re.match(r"^(\d+-)", ws.title.strip())
        if not m or m.group(1) not in SANTRALLER:
            continue
        kimlik, ad, canli, ozet_adi = SANTRALLER[m.group(1)]
        d = santral_sayfasi(ws)
        oz = ozet_bilgi.get(ozet_adi) or {}
        d["kw"] = oz.get("kw") or d["kw"]
        sonuc.append({
            "id": kimlik, "ad": ad, "durum": "Devrede",
            "canli": {"portal": canli[0], "ad": canli[1]} if canli else None,
            "yillik": oz.get("yillik", {}),
            **d,
        })

    # Son veri ayı: son yılın, herhangi bir santralde üretim görülen son ayı
    en_son_yil = max(y for s in sonuc for y in s["aylik"])
    son_ay = 0
    for ay in range(12):
        if any((s["aylik"].get(en_son_yil) or [None] * 12)[ay] not in (None, 0.0) for s in sonuc):
            son_ay = ay + 1
    for s in sonuc:  # son aydan sonraki (henüz girilmemiş) aylar boş sayılır
        if en_son_yil in s["aylik"]:
            s["aylik"][en_son_yil] = [v if i < son_ay else None for i, v in enumerate(s["aylik"][en_son_yil])]

    # Çapraz kontrol: aylık toplamlar yıllık özetle uyuşuyor mu? (Excel'deki tutarsızlıkları bildirir)
    for s in sonuc:
        for y, tutar in s["yillik"].items():
            aylik = s["aylik"].get(int(y))
            if aylik and tutar and int(y) < en_son_yil:
                toplam = sum(v for v in aylik if v)
                if abs(toplam - tutar) / tutar > 0.02:
                    notlar.append({"santral": s["ad"], "yil": int(y), "aylik_toplam": round(toplam),
                                   "ozet_sayfa": round(tutar), "fark_yuzde": round((toplam / tutar - 1) * 100, 1)})

    for s in sonuc:
        s["aylik"] = {str(y): [yuvarla(v, 1) for v in vals] for y, vals in sorted(s["aylik"].items())}
        s["beklenen_aylik"] = [yuvarla(v, 1) for v in s["beklenen_aylik"]]
    return sonuc, planlanan, {"yil": en_son_yil, "ay": son_ay}, notlar


# Arıza takip listesindeki ad -> santral kimliği (listedeki adlar sayfa adlarından farklı yazılıyor)
ARIZA_ADLARI = {
    "seyrek hayvan hastanesi": "seyrek-hayvan-hastanesi", "buca toros itfaiye": "buca-toros-itfaiye",
    "ekrem akurgal": "ekrem-akurgal", "örnekköy sosyal projeler": "ornekkoy-sosyal-projeler",
    "buz pisti havuz": "bornova-havuz", "selçuk transfer istasyonu": "selcuk-transfer",
    "mustafa necati kültür merkezi": "yesilyurt-kultur-merkezi", "bergama mezbahane": "bergama-mezbahasi",
    "aliağa itfaiye": "aliaga-itfaiye", "konak tünelleri": "konak-tuneller", "bergama itfaiye": "bergama-itfaiye",
    "uzundere spor kompleksi": "uzundere-spor", "aliağa kültür merkezi": "aliaga-kultur-merkezi",
    "seyrek itfaiye": "seyrek-itfaiye", "bornova kültür merkezi": "bornova-kultur-merkezi",
    "çiğli aile danışma merkezi": "cigli-aile-danisma", "gümüldür itfaiye": "gumuldur-itfaiye",
    "selçuk itfaiye": "selcuk-itfaiye", "çeşme itfaiye": "cesme-itfaiye", "torbalı çaybaşı itfaiye": "caybasi-itfaiye",
    "akkaya telsiz istasyonu dges": "akkaya-dges",
}


def ariza_dosyasi(yol):
    wb = openpyxl.load_workbook(yol, data_only=True)
    ws = wb.worksheets[0]
    donem = None
    for row in ws.iter_rows(min_row=1, max_row=3):
        for c in row:
            if isinstance(c.value, str) and re.search(r"\(\d{4}\)", c.value):
                donem = c.value.strip()
    satirlar = []
    for r in range(4, ws.max_row + 1):
        ad = ws.cell(r, 2).value
        if not ad:
            continue
        ad = str(ad).strip()
        if ad.upper() == "TOPLAM":
            continue
        kimlik = next((v for k, v in ARIZA_ADLARI.items() if tr_lower(ad).startswith(k)), None)
        satirlar.append({
            "id": kimlik,
            "ad": ad,
            "hataya_dusme": int(sayi(ws.cell(r, 3).value) or 0),
            "bekleme": int(sayi(ws.cell(r, 4).value) or 0),
            "izleme_disi": int(sayi(ws.cell(r, 5).value) or 0),
            "uretimsiz_gun": int(sayi(ws.cell(r, 6).value) or 0),
            "tarih": (str(ws.cell(r, 7).value).strip() if ws.cell(r, 7).value else None),
            "neden": (str(ws.cell(r, 8).value).strip() if ws.cell(r, 8).value else None),
        })
    return {"donem": donem, "satirlar": satirlar}


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    santraller, planlanan, son, notlar = uretim_dosyasi(sys.argv[1])
    cikti = {
        "olusturma": datetime.now().strftime("%Y-%m-%d"),
        "son_donem": son,
        "santraller": santraller,
        "planlanan": planlanan,
        "veri_notlari": notlar,
        "ariza": ariza_dosyasi(sys.argv[2]) if len(sys.argv) > 2 else None,
    }
    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(cikti, f, ensure_ascii=False, separators=(",", ":"))
    print(f"{len(santraller)} santral, {len(planlanan)} planlanan, son dönem {son['yil']}-{son['ay']:02d} -> {os.path.abspath(OUT_PATH)}")
    for n in notlar:
        print(f"UYARI: {n['santral']} {n['yil']}: aylık toplam {n['aylik_toplam']:,} kWh, özet sayfa {n['ozet_sayfa']:,} kWh (%{n['fark_yuzde']:+})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
