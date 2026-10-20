"""ISO 50001 enerji yönetimi verisinden data/iso50001.enc.json üretir (parolayla şifreli).

Kullanım:
    ELEKTRIK_PAROLA="parola" python collector/iso50001_to_json.py "<TSE GİDEN EVRAK klasörü>" ["<FRM.06.01 Risk ve Fırsat.xlsx>" ["<Verilerim.xlsx>"]]

Klasörde şunlar aranır:
  FRM.06.07 Enerji Referans Çizgileri/<yıl>/<no>.EnRÇ ... .xlsx   (her enerji kullanımı için aylık takip ve referans çizgisi)
  PLN.06.01 Enerji Eylem Planı.xlsx                                (tesis bazında eylemler)
İkinci argüman verilirse FRM.06.01 Risk ve Fırsat Değerlendirmesi de eklenir.

Depo herkese açık olduğu için düz metin YAZILMAZ; elektrik verisiyle aynı yöntemle (AES-256-GCM) şifrelenir.
Yatırım tutarları (TL) bilerek alınmaz. Excel dosyaları depoya konmaz.
"""
import datetime
import getpass
import glob
import json
import os
import re
import sys

import openpyxl

from elektrik_to_json import sifrele

OUT_PATH = os.environ.get("ISO_CIKTI") or os.path.join(os.path.dirname(__file__), "..", "data", "iso50001.enc.json")

TESISLER = {
    "aassm": "Ahmed Adnan Saygın Sanat Merkezi (AASSM)",
    "buz": "Buz Sporları Salonu",
    "buca": "Buca Sosyal Yaşam Kampüsü",
}

# dosya numarası -> (tesis, bölüm, enerji, birim)
SERILER = {
    1: ("aassm", "Ana bina", "Elektrik", "kWh"),
    2: ("aassm", "Ana bina", "Doğalgaz", "Sm³"),
    3: ("buz", "Buz Sporları Salonu", "Elektrik", "kWh"),
    4: ("buz", "Buz Sporları Salonu", "Doğalgaz", "Sm³"),
    5: ("buca", "Huzurevi", "Elektrik", "kWh"),
    6: ("buca", "Huzurevi", "Doğalgaz", "Sm³"),
    7: ("buca", "Aşevi", "Elektrik", "kWh"),
    8: ("buca", "Aşevi – Kazan", "Doğalgaz", "Sm³"),
    9: ("buca", "Aşevi – Mutfak (ocak)", "Doğalgaz", "Sm³"),
    10: ("buca", "Kapalı Spor Salonu", "Doğalgaz", "Sm³"),
    11: ("buca", "Soğutma Sistemi", "Elektrik", "kWh"),
}
EYLEM_SAYFALARI = {"Genel": "genel", "Buca Sosyal Yaşam Kampüsü": "buca", "Buz Sporları Salonu": "buz", "AASSM": "aassm"}
DEGISKEN_DISI = re.compile(r"hdd|cdd|etkinlik|yemek|sayı|sayi|çorba|corba|yaşlı|yasli|konaklayan", re.I)


def kucuk(s):
    return str(s).replace("İ", "i").replace("I", "ı").lower()


def sayi(x):
    if isinstance(x, bool) or x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(str(x).strip().replace(",", "."))
    except ValueError:
        return None


def yuvarla(x, n=1):
    return None if x is None else round(x, n)


def metin(x):
    if x is None:
        return None
    if isinstance(x, (datetime.datetime, datetime.date)):
        return x.strftime("%Y-%m-%d")
    t = re.sub(r"\s+", " ", str(x)).strip()
    return None if t in ("", "_", "-", "***", "0") else t


def referans_dosyasi(yol, takip_yil):
    wb = openpyxl.load_workbook(yol, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=1, max_row=22, max_col=32, values_only=True))
    r5, r6 = rows[4], rows[5]
    ref = next((i for i, x in enumerate(r5) if x and i > 0 and ("referans çi" in kucuk(x) or "beklenen" in kucuk(x))), None)
    gerekce = next((i for i, x in enumerate(r5) if x and "gerekçe" in kucuk(x)), None)
    if ref is None:
        raise ValueError("referans sütunu bulunamadı: " + yol)
    adaylar = [i for i in range(1, ref) if r6[i] and str(r6[i]).strip().startswith(str(takip_yil)) and not DEGISKEN_DISI.search(str(r6[i]))]
    if not adaylar:
        raise ValueError("takip sütunu bulunamadı: " + yol)
    takip = max(adaylar)
    degiskenler = {}
    for i in range(takip + 1, ref):
        if r6[i]:
            ad = re.sub(r"\s+", " ", str(r6[i])).strip()
            degiskenler[ad] = [yuvarla(sayi(rows[6 + a][i]), 2) for a in range(12)]
    hedef = None
    for i, x in enumerate(rows[19]):
        if x and "hedef" in kucuk(x):
            hedef = sayi(rows[20][i])
    t, r, g = [], [], []
    for a in range(12):
        satir = rows[6 + a]
        v = sayi(satir[takip])
        t.append(None if v is None else round(v, 1))
        r.append(yuvarla(sayi(satir[ref])))
        g.append(metin(satir[gerekce]) if gerekce is not None else None)
    return {"takip": t, "referans": r, "gerekce": g, "degiskenler": degiskenler, "hedef": hedef, "_excel_toplam": sayi(rows[18][takip])}


def eylem_plani(yol):
    wb = openpyxl.load_workbook(yol, data_only=True, read_only=True)
    sonuc = {}
    for sayfa, tid in EYLEM_SAYFALARI.items():
        if sayfa not in wb.sheetnames:
            continue
        liste = []
        for r in list(wb[sayfa].iter_rows(min_row=5, max_col=16, values_only=True)):
            if r[0] is None or not metin(r[5]):
                continue
            liste.append({
                "no": metin(r[0]), "enerji": metin(r[1]), "kapsam": metin(r[2]), "planlama": metin(r[3]), "gerekce": metin(r[4]),
                "onlem": metin(r[5]), "termin": metin(r[7]), "sorumlu": metin(r[8]), "hedef": metin(r[9]),
                "geri_odeme": metin(r[11]), "gerceklesme_tarihi": metin(r[12]), "durum": metin(r[13]),
            })
        sonuc[tid] = liste
    return sonuc


RISK_SAYFALARI = {"Risk ve Fırsat Değerlendirmesi": "genel", "BSYK": "buca", "AASSM": "aassm", "Buz Pisti": "buz"}


def risk_firsat(yol):
    """FRM.06.01 Risk ve Fırsat Değerlendirmesi: sayfa başına satırlar. No ve proses birleşik hücreden geldiği için aşağı taşınır."""
    wb = openpyxl.load_workbook(yol, data_only=True, read_only=True)
    sonuc = {}
    for sayfa, tid in RISK_SAYFALARI.items():
        if sayfa not in wb.sheetnames:
            continue
        liste, no, proses = [], None, None
        for r in wb[sayfa].iter_rows(min_row=7, max_col=14, values_only=True):
            if metin(r[1]):
                no = metin(r[1])
            if metin(r[2]):
                proses = metin(r[2])
            tur, tanim = metin(r[3]), metin(r[5])
            if tur not in ("Risk", "Fırsat") or not tanim:
                continue
            o, e = sayi(r[6]), sayi(r[7])
            seviye = sayi(r[8]) if sayi(r[8]) is not None else (o * e if o is not None and e is not None else None)
            liste.append({"no": no, "proses": proses, "tur": tur, "kaynak": metin(r[4]), "tanim": tanim, "o": o, "e": e, "seviye": seviye,
                          "sonra": sayi(r[9]), "aksiyon": metin(r[10]), "sorumlu": metin(r[11]), "faaliyet": metin(r[12])})
        sonuc[tid] = liste
    return sonuc


AYLAR_TR = ["ocak", "şubat", "mart", "nisan", "mayıs", "haziran", "temmuz", "ağustos", "eylül", "ekim", "kasım", "aralık"]
# seri no -> (kaynak türü, kaynak adı) ; 2025-2026 gerçekleşen tüketim Verilerim.xlsx'ten alınır
VERILERIM = {
    1: ("fark", "AASSM"),                      # AASSM sayfası: toplam − mahsup edilen
    2: ("ozet", "AASSM"),                      # Ayrıştırılmış Tüketimler özet matrisi
    3: ("havuz", None),                        # Buz Pisti-Havuz Ayırma: havuz çıkarılmış
    4: ("ozet", "Buz pisti"),
    5: ("ham_el", "HUZUREVİ ELEKTRİK (795)"),
    6: ("ham_gaz", "HUZUREVİ DOĞALGAZ"),
    7: ("fark", "Buca Aşevi"),                 # Buca Aşevi sayfası: toplam − Necati Kültür Merkezi
    8: ("ozet", "AŞEVİ KAZAN"),
    9: ("ozet", "AŞEVİ FIRIN"),
    10: ("ozet", "Kapalı spor salonu"),
    11: ("ham_el", "İKLİMLENDİRME ELEKTRİK (855)"),
}
# referans yöntemi: ortalama = (2022+2023)/2 × (1−hedef), 2024 dosyasındaki referansla aynı (sabit baz dönem);
# regresyon = (kesişim + katsayı × HDD/CDD) × (1−hedef), katsayılar 2024 EnRÇ dosyasındaki regresyon çıktısından
REGRESYON = {2, 6, 8, 10, 11}
# Verilerim.xlsx'in DEĞİŞKENLER sayfası 2025'te bitiyor; eksik yılların derece-gün değerleri MGM'den
# (mgm.gov.tr/Veridegerlendirme/gun-derece.aspx, İzmir merkez 35-00). Aynı tanım: 2024-2025 değerleri Excel'dekiyle birebir tutuyor.
# Aylar Ocak..Aralık; henüz açıklanmamış aylar None. Sayfada o yıl varsa sayfa esas alınır.
MGM_DERECE_GUN = {
    2026: {"HDD": [216, 162, 190, 40, 17, 0, 0, 0, 0, None, None, None],
           "CDD": [0, 0, 0, 0, 16, 153, 221, 234, 90, None, None, None]},  # Eylül kısmi (22 gün)
}


def aylik12():
    return [None] * 12


def ay_no(ad):
    return AYLAR_TR.index(kucuk(str(ad).strip()))


def regresyon_katsayilari(yol):
    wb = openpyxl.load_workbook(yol, data_only=True, read_only=True)
    satirlar = list(wb[wb.sheetnames[0]].iter_rows(min_row=1, max_row=60, max_col=8, values_only=True))
    for i, r in enumerate(satirlar):
        if r[0] and str(r[0]).strip() == "Kesişim":
            a = next(sayi(x) for x in r[1:] if sayi(x) is not None)
            b = next(sayi(x) for x in satirlar[i + 1][1:] if sayi(x) is not None)
            return a, b, str(satirlar[i + 1][0]).strip().upper()
    raise ValueError("regresyon katsayısı bulunamadı: " + yol)


def verilerim_oku(yol, klasor, seriler, notlar):
    wb = openpyxl.load_workbook(yol, data_only=True, read_only=True)
    # Ham Veri: abone × yıl × ay
    el, gaz, son = {}, {}, (0, 0)
    for r in wb["Ham Veri"].iter_rows(min_row=6, max_col=12, values_only=True):
        if r[1] and sayi(r[2]) and r[3] and sayi(r[4]) is not None:
            el.setdefault(str(r[1]).strip(), {}).setdefault(int(sayi(r[2])), aylik12())[ay_no(r[3])] = sayi(r[4])
            son = max(son, (int(sayi(r[2])), ay_no(r[3]) + 1))
        if r[8] and sayi(r[9]) and r[10] and sayi(r[11]) is not None:
            gaz.setdefault(str(r[8]).strip(), {}).setdefault(int(sayi(r[9])), aylik12())[ay_no(r[10])] = sayi(r[11])
    # Ayrıştırılmış Tüketimler: 73. satırdan sonra bölüm başına yıl × 12 ay
    ozet, etiket = {}, None
    for r in wb["Ayrıştırılmış Tüketimler"].iter_rows(min_row=73, max_col=16, values_only=True):
        if r[2]:
            etiket = str(r[2]).strip()
        if etiket and sayi(r[3]) and 2000 < sayi(r[3]) < 2100 and any(isinstance(x, (int, float)) for x in r[4:16]):
            ozet.setdefault(etiket, {})[int(sayi(r[3]))] = [sayi(x) for x in r[4:16]]

    def fark(sayfa):
        sh = list(wb[sayfa].iter_rows(min_row=4, max_row=15, max_col=11, values_only=True))
        sonuc = {}
        for k, y in enumerate((2023, 2024, 2025, 2026)):
            sonuc[y] = [None if sayi(r[1 + k]) is None else sayi(r[1 + k]) - (sayi(r[7 + k]) or 0) for r in sh]
        return sonuc

    havuz, buz = {}, {}
    for r in wb["Buz Pisti-Havuz Ayırma"].iter_rows(min_row=4, max_col=4, values_only=True):
        if isinstance(r[0], datetime.datetime) and sayi(r[3]) is not None:
            y, a = r[0].year, r[0].month - 1
            buz.setdefault(y, aylik12())[a] = sayi(r[3])
            havuz.setdefault(y, {"toplam": aylik12(), "havuz": aylik12()})
            havuz[y]["toplam"][a] = sayi(r[1]); havuz[y]["havuz"][a] = sayi(r[2])
    # HDD / CDD
    hdd, cdd = {}, {}
    for r in wb["DEĞİŞKENLER Ort."].iter_rows(min_row=5, max_col=4, values_only=True):
        if isinstance(r[0], datetime.datetime):
            hdd.setdefault(r[0].year, aylik12())[r[0].month - 1] = sayi(r[3])
            cdd.setdefault(r[0].year, aylik12())[r[0].month - 1] = sayi(r[1])
    for y, d in MGM_DERECE_GUN.items():
        hdd.setdefault(y, d["HDD"])
        cdd.setdefault(y, d["CDD"])

    for no, s in seriler.items():
        tur, ad = VERILERIM[no]
        deger = {"fark": lambda: fark(ad), "ozet": lambda: ozet.get(ad, {}), "havuz": lambda: buz,
                 "ham_el": lambda: el.get(ad, {}), "ham_gaz": lambda: gaz.get(ad, {})}[tur]()
        # doğrulama: 2023-2024 değerleri eski EnRÇ dosyalarıyla aynı olmalı
        for y in ("2023", "2024"):
            if y in s["yillar"] and int(y) in deger:
                ben = sum(x for x in deger[int(y)] if x is not None)
                eski = sum(x for x in s["yillar"][y]["takip"] if x is not None)
                if abs(ben - eski) > 1:
                    notlar.append({"seri": f"{s['tesis']} / {s['bolum']} / {s['enerji']}", "yil": y, "excel_toplam": round(eski), "aylik_toplam": round(ben),
                                   "not": "Verilerim.xlsx toplamı EnRÇ dosyasından farklı (EnRÇ değeri kullanıldı)."})
        ref24 = s["yillar"]["2024"]
        hedef = ref24["hedef"] or 0.02
        reg = None
        if no in REGRESYON:
            dosya = [f for f in glob.glob(os.path.join(klasor, "2024", f"{no}.*.xlsx")) if not os.path.basename(f).startswith("~$")][0]
            a, b, degisken = regresyon_katsayilari(dosya)
            reg = (a, b, degisken)
        for y in (2025, 2026):
            if y not in deger:
                continue
            takip = [None if (v is None or (y, i + 1) > son) else round(v, 1) for i, v in enumerate(deger[y])]
            if reg:
                kaynak = hdd if reg[2] == "HDD" else cdd
                var = kaynak.get(y)
                ref = [None if (var is None or var[i] is None) else round((reg[0] + reg[1] * var[i]) * (1 - hedef), 1) for i in range(12)]
                degisken = {f"{reg[2]} {y}": var} if var else {}
            else:
                ref, degisken = list(ref24["referans"]), {}
            s["yillar"][str(y)] = {"takip": takip, "referans": ref, "gerekce": [None] * 12, "degiskenler": degisken, "hedef": hedef,
                                   "yontem": "regresyon" if reg else "ortalama"}
        if reg and 2026 in deger and 2026 not in (hdd if reg[2] == "HDD" else cdd):
            notlar.append({"seri": f"{s['tesis']} / {s['bolum']} / {s['enerji']}", "yil": "2026", "excel_toplam": None, "aylik_toplam": None,
                           "not": f"2026 için {reg[2]} değerleri Verilerim.xlsx'te yok; referans çizgisi hesaplanamadı (yalnızca gerçekleşen gösteriliyor)."})
    motorin = {}
    for r in wb["Motorin"].iter_rows(min_row=4, max_row=12, max_col=6, values_only=True):
        if sayi(r[2]) and 2000 < sayi(r[2]) < 2100:
            for k, t in ((3, "buca"), (4, "aassm"), (5, "buz")):
                if sayi(r[k]) is not None:
                    motorin.setdefault(t, {})[str(int(sayi(r[2])))] = sayi(r[k])
    return {"motorin": motorin, "havuz": {str(y): v for y, v in havuz.items()}, "son_donem": {"yil": son[0], "ay": son[1]}}


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    kok = sys.argv[1]
    klasor = os.path.join(kok, "FRM.06.07 Enerji Referans Çizgileri")
    seriler = {}
    for yil in (2023, 2024):
        for yol in sorted(glob.glob(os.path.join(klasor, str(yil), "*.xlsx"))):
            ad = os.path.basename(yol)
            m = re.match(r"(\d+)\.", ad)
            if ad.startswith("~$") or not m or int(m.group(1)) not in SERILER:
                continue
            no = int(m.group(1))
            tesis, bolum, enerji, birim = SERILER[no]
            s = seriler.setdefault(no, {"id": f"{tesis}-{no}", "tesis": tesis, "bolum": bolum, "enerji": enerji, "birim": birim, "yillar": {}})
            s["yillar"][str(yil)] = referans_dosyasi(yol, yil)
            print(f"{yil} {ad}: ok")
    notlar = []
    for s in seriler.values():
        for yil, v in s["yillar"].items():
            excel = v.pop("_excel_toplam")
            ben = sum(x for x in v["takip"] if x is not None)
            if excel is None or abs(excel - ben) > 1:
                notlar.append({"seri": f"{s['tesis']} / {s['bolum']} / {s['enerji']}", "yil": yil, "excel_toplam": None if excel is None else round(excel), "aylik_toplam": round(ben),
                               "not": "Excel'deki toplam satırı aylık değerlerin toplamıyla uyuşmuyor (formül bir ayı dışarıda bırakıyor olabilir); aylık değerlerin toplamı kullanıldı."})
    eylem = eylem_plani(os.path.join(kok, "PLN.06.01 Enerji Eylem Planı.xlsx"))
    risk = risk_firsat(sys.argv[2]) if len(sys.argv) > 2 else {}
    ek = {}
    if len(sys.argv) > 3:
        ek = verilerim_oku(sys.argv[3], klasor, seriler, notlar)
    cikti = {
        "olusturma": datetime.date.today().isoformat(),
        "tesisler": [{"id": k, "ad": v} for k, v in TESISLER.items()],
        "seriler": [seriler[k] for k in sorted(seriler)],
        "eylem": eylem,
        "risk": risk,
        "ek": ek,
        "veri_notlari": notlar,
    }
    parola = os.environ.get("ELEKTRIK_PAROLA") or getpass.getpass("Parola: ")
    if len(parola) < 12:
        print("Parola en az 12 karakter olmalı.")
        return 1
    os.makedirs(os.path.dirname(os.path.abspath(OUT_PATH)), exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(sifrele(cikti, parola), f, separators=(",", ":"))
    print(f"{len(seriler)} seri, eylem: { {k: len(v) for k, v in eylem.items()} } -> {os.path.abspath(OUT_PATH)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
