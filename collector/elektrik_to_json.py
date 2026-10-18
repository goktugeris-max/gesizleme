"""Elektrik tüketimi Excel'inden data/elektrik.json üretir.

Kullanım:
    python collector/elektrik_to_json.py "Elektrik.xlsx"

Depo ve site herkese açık olduğu için ABONE NO, ADRES, kabin bilgisi ve fatura TUTARLARI (TL)
çıktıya YAZILMAZ; yalnızca toplamlar (ay / daire başkanlığı / tarife / tedarikçi, kWh) yazılır.

Kaynaklar (Excel sayfaları):
  Veri          ham aylık abone tüketimi (2024 → ...): DB, tarife, tedarikçi için toplamlar
  Yıllık Durum  aylık faturalı tüketim (2023 → ...), genel aydınlatma hariç
  Sonuç         genel aydınlatmadan gelen %20 pay, müstecir noktalar, GES net üretim, üretim fazlası
"""
import collections
import json
import os
import sys
from datetime import datetime

import openpyxl

OUT_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "elektrik.json")
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]


def sayi(x):
    if isinstance(x, bool) or x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(str(x).strip().replace(",", "."))
    except ValueError:
        return None


def yuvarla(x, n=1):
    return None if x is None else round(x, n)


def ham_veri(wb):
    """Veri sayfası -> toplamlar. Abone numarası yalnızca ay başına abone SAYISI için kullanılır."""
    ws = wb["Veri"]
    db = collections.defaultdict(lambda: collections.defaultdict(lambda: [0.0] * 12))
    tarife = collections.defaultdict(lambda: collections.defaultdict(lambda: [0.0] * 12))
    tedarikci = collections.defaultdict(lambda: collections.defaultdict(lambda: [0.0] * 12))
    toplam = collections.defaultdict(lambda: [0.0] * 12)
    aboneler = collections.defaultdict(set)
    sayac = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[1] is None:
            continue
        yil, ay = sayi(r[2]), sayi(r[3])
        kwh = sayi(r[7])
        if yil is None or ay is None or kwh is None:
            continue
        yil, ay = int(yil), int(ay)
        i = ay - 1
        ad = str(r[4]).strip() if r[4] not in (None, 0, "0") else "Belirtilmemiş"
        t = str(r[5]).strip() if r[5] not in (None, 0, "0") else "Belirtilmemiş"
        s = str(r[6]).strip() if r[6] not in (None, 0, "0") else "Belirtilmemiş"
        db[ad][yil][i] += kwh
        tarife[t][yil][i] += kwh
        tedarikci[s][yil][i] += kwh
        toplam[yil][i] += kwh
        aboneler[(yil, ay)].add(str(r[1]))
        sayac += 1
    abone_sayisi = collections.defaultdict(lambda: [None] * 12)
    for (yil, ay), kume in aboneler.items():
        abone_sayisi[yil][ay - 1] = len(kume)
    return db, tarife, tedarikci, toplam, abone_sayisi, sayac


def yillik_durum(wb):
    """Yıllık Durum sayfası: ay satırları 3..14, yıl sütunları 2023..2026 (B..E)."""
    ws = wb["Yıllık Durum"]
    yillar = {c: ws.cell(2, c).value for c in range(2, 6)}
    sonuc = {}
    for c, y in yillar.items():
        if y is None:
            continue
        sonuc[int(y)] = [sayi(ws.cell(2 + a, c).value) for a in range(1, 13)]
    return sonuc


def sonuc_sayfasi(wb):
    """Sonuç sayfası: her yıl bloğu için E (aydınlatma %20 pay), F (müstecir), G (GES net), H (fazla)."""
    ws = wb["Sonuç"]
    bloklar = {}
    r = 1
    while r <= ws.max_row:
        yil = ws.cell(r, 2).value
        if sayi(yil) is not None and 2015 <= int(sayi(yil)) <= 2100 and ws.cell(r, 3).value in (None, "") and ws.cell(r, 4).value:
            y = int(sayi(yil))
            blok = {k: [] for k in ("faturali", "aydinlatma", "mustecir", "ges_net", "fazla")}
            for a in range(12):
                row = r + 1 + a
                for anahtar, c in (("faturali", 4), ("aydinlatma", 5), ("mustecir", 6), ("ges_net", 7), ("fazla", 8)):
                    blok[anahtar].append(sayi(ws.cell(row, c).value))
            bloklar[y] = blok
            r += 13
        else:
            r += 1
    return bloklar


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True, read_only=True)
    db, tarife, tedarikci, toplam, abone_sayisi, satir_sayisi = ham_veri(wb)
    yd = yillik_durum(wb)
    sn = sonuc_sayfasi(wb)
    notlar = []

    # Son veri ayı: en son yılda faturalı tüketimi 0'dan büyük son ay
    en_son_yil = max(yd)
    son_ay = max((i + 1 for i, v in enumerate(yd[en_son_yil]) if v), default=0)

    aylik = {}
    for yil in sorted(yd):
        fat = yd[yil]
        blok = sn.get(yil, {})
        yeni = {k: [] for k in ("faturali", "aydinlatma_pay", "mustecir", "ges_net", "fazla", "gercek")}
        for i in range(12):
            bos = yil == en_son_yil and i >= son_ay
            f = None if bos else fat[i]
            e = None if bos else (blok.get("aydinlatma") or [None] * 12)[i]
            m = None if bos else (blok.get("mustecir") or [None] * 12)[i]
            g = None if bos else (blok.get("ges_net") or [None] * 12)[i]
            h = None if bos else (blok.get("fazla") or [None] * 12)[i]
            # Excel'deki formül: Gerçek = Faturalı + Aydınlatma payı − Müstecir − Fazla + GES net
            gercek = None
            if f is not None and None not in (e, m, g, h):
                gercek = f + e - m - h + g
            for k, v in zip(("faturali", "aydinlatma_pay", "mustecir", "ges_net", "fazla", "gercek"), (f, e, m, g, h, gercek)):
                yeni[k].append(yuvarla(v))
            # Sonuç sayfasındaki faturalı değer, Yıllık Durum'dan farklıysa not et
            sf = (blok.get("faturali") or [None] * 12)[i]
            if not bos and sf is not None and f is not None and abs(sf - f) > 1:
                notlar.append({"yil": yil, "ay": AYLAR[i], "sonuc_sayfasi": round(sf), "yillik_durum": round(f), "fark": round(f - sf)})
        aylik[str(yil)] = yeni

    def kisalt(sozluk):
        return {ad: {str(y): [yuvarla(v) for v in vals] for y, vals in sorted(yillar.items())} for ad, yillar in sorted(sozluk.items())}

    # Ham veri toplamı Yıllık Durum ile uyuşuyor mu?
    for y, vals in toplam.items():
        for i, v in enumerate(vals):
            ref = (yd.get(y) or [None] * 12)[i]
            if ref and abs(v - ref) > 1:
                notlar.append({"yil": y, "ay": AYLAR[i], "ham_veri_toplami": round(v), "yillik_durum": round(ref), "fark": round(ref - v)})

    cikti = {
        "olusturma": datetime.now().strftime("%Y-%m-%d"),
        "son_donem": {"yil": en_son_yil, "ay": son_ay},
        "aylik": aylik,
        "abone_sayisi": {str(y): v for y, v in sorted(abone_sayisi.items())},
        "db": kisalt(db),
        "tarife": kisalt(tarife),
        "tedarikci": kisalt(tedarikci),
        "veri_notlari": notlar,
    }
    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(cikti, f, ensure_ascii=False, separators=(",", ":"))
    print(f"{satir_sayisi} ham satır, {len(db)} daire başkanlığı, son dönem {en_son_yil}-{son_ay:02d} -> {os.path.abspath(OUT_PATH)}")
    for n in notlar:
        print("NOT:", n)
    return 0


if __name__ == "__main__":
    sys.exit(main())
