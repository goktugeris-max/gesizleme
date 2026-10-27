#!/usr/bin/env python3
"""ISO 50001 EnYS belgelerinden panel için data/enys.json üretir (ŞİFRESİZ, herkese açık — kullanıcı kararı).

Kullanım:  python3 collector/enys_to_json.py "<EnYS_Baglam klasörü>"
Okunan dosyalar: FRM.06.02 (amaç/hedefler), FRM.06.05 (EnPG), FRM.06.04 (ÖEK değerlendirme), FRM.06.06 (ölçme-izleme planı),
LST.09.01 (iç tetkik soru listesi), 12/13.EnR* (Eşrefpaşa Semt Polikliniği referans çizgileri).
YAZILMAZ: sayaç/tesisat numaraları, OSOS bilgileri, kişi adları, toplantı tutanağı dosya adları. Excel/PDF dosyaları depoya konmaz.
Eşrefpaşa serileri Buca Sosyal Yaşam Kampüsü'nün bölümü olarak (buca-12/13) ISO şifreli verisine tarayıcıda eklenir (iso.js → Enys.birlestir).
"""
import glob
import json
import os
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import iso50001_to_json as ISO  # referans_dosyasi, sayi, metin

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "enys.json")


def hucre(x):
    if x is None:
        return None
    t = re.sub(r"\s+", " ", str(x)).strip()
    return t or None


def satirlar(yol, sayfa, **kw):
    wb = openpyxl.load_workbook(yol, data_only=True)
    return list(wb[sayfa].iter_rows(values_only=True, **kw)), wb.sheetnames


# ---------------- hedefler (FRM.06.02) ----------------
KAPSAM = [  # (amaçta aranan, tesis, bölüm öneki)
    ("eşrefpaşa", "buca", "Eşrefpaşa Semt Polikliniği"), ("huzurevi", "buca", "Huzurevi"), ("aşevi", "buca", "Aşevi"),
    ("kapalı spor", "buca", "Kapalı Spor Salonu"), ("soğutma", "buca", "Soğutma Sistemi"),
    ("buz sporları", "buz", None), ("aassm", "aassm", None),
]


def hedef_coz(amac, hedef):
    a, h = amac.lower().replace("İ", "i"), hedef.lower().replace("İ", "i")
    kap = next((k for k in KAPSAM if k[0] in a), None)
    baz = re.search(r"(\d{4}) y[ıi]l[ıi] ayl[ıi]k t[üu]ketimlerine g[öo]re", h)
    if not kap or not baz:
        return {"tur": "nitel"}
    ener = [e for e, k in (("Elektrik", "elektrik"), ("Doğalgaz", "doğalgaz")) if k in a]
    oranlar = {}
    e = re.search(r"elektrik t[üu]ketimi hedefi %\s*(\d+)", h)
    g = re.search(r"do[ğg]algaz t[üu]ketimi hedefi %\s*(\d+)", h)
    if e or g:
        if e:
            oranlar["Elektrik"] = int(e.group(1)) / 100
        if g:
            oranlar["Doğalgaz"] = int(g.group(1)) / 100
    else:
        t = re.search(r"%\s*(\d+)", h)
        if not t:
            return {"tur": "nitel"}
        for x in ener or ["Elektrik", "Doğalgaz"]:
            oranlar[x] = int(t.group(1)) / 100
    return {"tur": "nicel", "baz_yil": int(baz.group(1)), "tesis": kap[1], "bolum": kap[2],
            "kalemler": [{"enerji": k, "oran": v} for k, v in oranlar.items()]}


def hedefler(yol):
    wb = openpyxl.load_workbook(yol, data_only=True)
    sonuc = {}
    for sh in wb.sheetnames:
        liste = []
        for r in wb[sh].iter_rows(values_only=True):
            if r[0] is None or not str(r[0]).strip().isdigit() or not hucre(r[1]):
                continue
            amac, hedef = hucre(r[1]), hucre(r[2]) or ""
            liste.append({"sira": int(r[0]), "amac": amac, "hedef": hedef, "eylemler": hucre(r[3]), "sorumlu": hucre(r[4]), **hedef_coz(amac, hedef)})
        if liste:
            sonuc[sh] = liste
    return sonuc


# ---------------- EnPG (FRM.06.05) ----------------
def enpg(yol):
    wb = openpyxl.load_workbook(yol, data_only=True)
    sonuc = {}
    for sh in wb.sheetnames:
        satir = list(wb[sh].iter_rows(values_only=True))
        bas = next((i for i, r in enumerate(satir) if r[0] and str(r[0]).strip() == "No"), None)
        if bas is None:
            continue
        kol = {ISO.kucuk(str(v)): j for j, v in enumerate(satir[bas]) if v}
        al = lambda r, *anah: next((r[j] for k, j in kol.items() if all(a in k for a in anah)), None)
        yil = re.search(r"(\d{4})", sh).group(1)
        bina, liste = None, []
        for r in satir[bas + 1:]:
            if r[0] is None or not str(r[0]).strip().isdigit():
                continue
            bina = hucre(al(r, "bina")) or bina
            kwh = ISO.sayi(al(r, "yıllık toplam"))
            deger = ISO.sayi(al(r, "enpg değeri"))
            if kwh is None and deger is None:
                continue
            liste.append({"no": int(r[0]), "bina": re.sub(r"^Buca Sosyal Yaşam Kampüsü\s*\(?|\)$", "", bina or "").strip() or bina, "kullanim": hucre(al(r, "enerji kullanımı")),
                          "tur": hucre(al(r, "enerji türü")), "kwh": kwh and round(kwh, 1), "sabit": ISO.sayi(al(r, "faaliyet sabiti")),
                          "degisken": ISO.sayi(al(r, "faaliyet değişkeni")), "tanim": hucre(al(r, "enpg tanımı")), "deger": deger and round(deger, 4)})
        sonuc[yil] = liste
    return sonuc


# ---------------- ÖEK değerlendirme ve ölçme planı (sayaç numaraları ALINMAZ) ----------------
def oek(yol_04, yol_06):
    out = {}
    try:
        s04 = list(openpyxl.load_workbook(yol_04, data_only=True)["2025"].iter_rows(values_only=True))
        bas = next(i for i, r in enumerate(s04) if r[0] and str(r[0]).strip() == "No")
        for r in s04[bas + 1:]:
            if r[0] is not None and str(r[0]).strip().isdigit():
                out[int(r[0])] = {"oran_pct": ISO.sayi(r[5]) if len(r) > 5 else None, "degiskenler": hucre(r[8]) if len(r) > 8 else None, "eylem": hucre(r[12]) if len(r) > 12 else None}
    except Exception as e:
        print("uyarı: FRM.06.04 okunamadı:", e)
    try:
        s06 = list(openpyxl.load_workbook(yol_06, data_only=True)["2026"].iter_rows(values_only=True))
        bas = next(i for i, r in enumerate(s06) if r[0] and str(r[0]).strip() == "No")
        baslik = [ISO.kucuk(str(v)) if v else "" for v in s06[bas - 1]]
        for r in s06[bas + 1:]:
            if r[0] is not None and str(r[0]).strip().isdigit():
                d = out.setdefault(int(r[0]), {})
                d["degisken_izleme"] = hucre(r[5])
                d["okuma"] = hucre(r[6])
                d["periyot"] = hucre(r[9])          # r[8] = sayaç numarası: ALINMAZ
                d["kayit_sorumlusu"] = hucre(r[10])
    except Exception as e:
        print("uyarı: FRM.06.06 okunamadı:", e)
    return out


# ---------------- iç tetkik soruları (LST.09.01) ----------------
def tetkik(yol):
    ws = openpyxl.load_workbook(yol, data_only=True).active
    satir = list(ws.iter_rows(values_only=True))
    bas = next(i for i, r in enumerate(satir) if r[0] and "KURULU" in str(r[0]).upper() and any(r[1:]))
    isimler = [(j, hucre(v)) for j, v in enumerate(satir[bas]) if j >= 2 and v]
    sorular_bas = next(i for i, r in enumerate(satir) if hucre(r[0]) == "SORULAR")
    gruplar, madde = [], None
    for r in satir[sorular_bas + 1:]:
        t = hucre(r[0])
        if not t:
            continue
        if re.match(r"^\d+\.\s*\d+", t) or re.match(r"^\d+\.\s*\D", t):          # "4.1 ...", "6.Planlama"
            m = re.match(r"^(\d+(?:\.\s*\d+)*)\.?\s*(.*)$", t)
            madde = {"madde": re.sub(r"\s+", "", m.group(1)), "baslik": m.group(2), "sorular": []}
            gruplar.append(madde)
            continue
        if madde is None:
            continue
        kimler = [ad for j, ad in isimler if j < len(r) and r[j] and str(r[j]).strip().lower() == "x"]
        no = re.match(r"^(\d+)\s*[-•.]\s*(.*)$", t, re.S)
        madde["sorular"].append({"no": int(no.group(1)) if no else None, "soru": (no.group(2) if no else t).strip(), "kanit": hucre(r[1]), "kimler": kimler})
    return [g for g in gruplar if g["sorular"]]


# ---------------- Eşrefpaşa serileri ----------------
def esrefpasa(klasor):
    seriler = []
    for no, enerji, birim in ((12, "Elektrik", "kWh"), (13, "Doğalgaz", "Sm³")):
        yol = [f for f in glob.glob(os.path.join(klasor, f"{no}.*.xlsx")) if not os.path.basename(f).startswith("~$")][0]
        r26 = ISO.referans_dosyasi(yol, 2026)
        r26.pop("_excel_toplam", None)
        r26["yontem"] = "ortalama"
        ws = openpyxl.load_workbook(yol, data_only=True)["Veriler"]
        yillar, yil = {"2024": [], "2025": []}, None
        for r in ws.iter_rows(min_row=1, max_row=40, max_col=4, values_only=True):   # B sütununda yıl etiketi, C sütununda aylık tüketim
            if r[1] in (2024, 2025):
                yil = str(int(r[1]))
            if yil and len(yillar[yil]) < 12 and ISO.sayi(r[2]) is not None:
                yillar[yil].append(round(ISO.sayi(r[2]), 1))
        assert len(yillar["2024"]) == 12 and len(yillar["2025"]) == 12, f"{no}: Veriler sayfası beklenen düzende değil"
        y = {"2024": {"takip": yillar["2024"], "referans": [None] * 12, "gerekce": [None] * 12, "degiskenler": {}, "hedef": None, "yontem": "ortalama"},
             "2025": {"takip": yillar["2025"], "referans": [None] * 12, "gerekce": [None] * 12, "degiskenler": {}, "hedef": None, "yontem": "ortalama"},
             "2026": r26}
        ort = [None if (a is None or b is None) else round((a + b) / 2, 1) for a, b in zip(yillar["2024"], yillar["2025"])]
        assert all(abs((o or 0) - (r or 0)) < 1 for o, r in zip(ort, r26["referans"])), f"{no}: referans 2024-2025 ortalaması tutmadı"
        seriler.append({"id": f"buca-{no}", "tesis": "buca", "bolum": "Eşrefpaşa Semt Polikliniği", "enerji": enerji, "birim": birim, "yillar": y})
    return seriler


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    k = sys.argv[1]
    g = lambda *p: glob.glob(os.path.join(k, *p))[0]
    cikti = {
        "kaynak": "İBB ISO 50001 EnYS belgeleri (FRM.06.02/04/05/06, LST.09.01, EnRÇ 2026)",
        "hedefler": hedefler(g("FRM.06.02*.xlsx")),
        "enpg": enpg(g("FRM.06.05*.xlsx")),
        "oek": oek(g("FRM.06.04*.xlsx"), g("FRM.06.06*.xlsx")),
        "tetkik": tetkik(g("LST.09.01*.xlsx")),
        "iso_ek": {"seriler": esrefpasa(k)},
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(cikti, f, ensure_ascii=False, separators=(",", ":"))
    print("hedef yılları:", {y: len(v) for y, v in cikti["hedefler"].items()}, "| nicel 2026:", sum(1 for h in cikti["hedefler"].get("2026", []) if h["tur"] == "nicel"))
    print("EnPG:", {y: len(v) for y, v in cikti["enpg"].items()}, "| ÖEK:", len(cikti["oek"]), "| tetkik:", len(cikti["tetkik"]), "madde,", sum(len(x["sorular"]) for x in cikti["tetkik"]), "soru")
    print("yazıldı:", os.path.normpath(OUT), os.path.getsize(OUT), "bayt")
    return 0


if __name__ == "__main__":
    sys.exit(main())
